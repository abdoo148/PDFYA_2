"""منطق البرنامج: توقيع PDF، تحويل الصيغ، أدوات PDF، مكتبة التواقيع. بدون واجهة (قابل للاختبار)."""
import datetime
import html
import io
import json
import os
import re
import shutil
import subprocess
import sys
import threading
import time
import unicodedata
import urllib.request

try:
    import pymupdf as fitz  # PyMuPDF
except ImportError:
    import fitz
from PIL import Image

from i18n import tr

APP_NAME = "PDFYA"
APP_TAGLINE = tr("موقّع ومحوّل الملفات")
PUBLISHER = "Codpact"
PUBLISHER_AR = tr("Codpact للحلول الرقمية")
LEGACY_APP_NAMES = ("Pedo", "SignConvert")  # أسماء سابقة: تُنسخ بيانات أحدثها تلقائيًا عند أول تشغيل
APP_VERSION = "1.5.0"

APP_LICENSE = "AGPL-3.0"
# رابط التحقق من التحديثات (فارغ = معطّل). يقبل:
#   - GitHub Releases: https://api.github.com/repos/<user>/<repo>/releases/latest
#   - أو ملف JSON بالشكل: {"version": "1.3.0", "url": "رابط صفحة التنزيل", "notes": "ما الجديد"}
# يمكن تغييره بدون إعادة بناء عبر المفتاح "update_url" في settings.json
UPDATE_URL = "https://api.github.com/repos/abdoo148/PDFYA_2/releases/latest"

# PyMuPDF لا يدعم العمل المتوازي: كل استدعاء له من خيط آخر يجب أن يمسك هذا القفل.
FITZ_LOCK = threading.RLock()


class Cancelled(Exception):
    """يُرفع من دالة التقدم عندما يضغط المستخدم إلغاء."""

    def __init__(self):
        super().__init__(tr("تم الإلغاء"))


def _tick(progress, done, total):
    """يبلّغ الواجهة بنسبة التقدم (وقد يرفع Cancelled)."""
    if progress:
        progress(done, total)


# ---------------------------------------------------------------- أدوات مساعدة
def parse_ranges(text, n_pages):
    """'1-3,5' -> [0,1,2,4] (فهارس تبدأ من صفر، بنفس الترتيب المكتوب). يرفع ValueError عند خطأ."""
    pages = []
    for part in text.replace("،", ",").split(","):
        part = part.strip()
        if not part:
            continue
        try:
            if "-" in part:
                a, b = part.split("-", 1)
                a, b = int(a), int(b)
            else:
                a = b = int(part)
        except ValueError:
            raise ValueError(tr("نطاق غير صحيح: {0}", part)) from None
        step = 1 if b >= a else -1
        if min(a, b) < 1 or max(a, b) > n_pages:
            raise ValueError(tr("الصفحات المسموحة من 1 إلى {0}", n_pages))
        pages.extend(range(a - 1, b - 1 + step, step))
    if not pages:
        raise ValueError(tr("لم يتم تحديد صفحات"))
    return pages


def unique_path(path):
    base, ext = os.path.splitext(path)
    i = 1
    while os.path.exists(path):
        path = f"{base} ({i}){ext}"
        i += 1
    return path


def size_text(path):
    s = os.path.getsize(path)
    return tr("{0:.0f} ك.ب", s / 1024) if s < 1024 * 1024 else tr("{0:.2f} م.ب", s / 1024 / 1024)


def open_pdf(path, password=None):
    doc = fitz.open(path)
    if doc.needs_pass and not (password and doc.authenticate(password)):
        doc.close()
        raise ValueError(tr("الملف محمي بكلمة مرور: {0}", os.path.basename(path)))
    return doc


def _save(doc, out_path, **kw):
    opts = dict(garbage=3, deflate=True)
    opts.update(kw)
    doc.save(out_path, **opts)


def data_dir():
    base = os.environ.get("APPDATA") or os.path.expanduser("~")
    d = os.path.join(base, APP_NAME)
    if not os.path.isdir(d):
        for legacy in LEGACY_APP_NAMES:  # نسخ مكتبة التواقيع والإعدادات من إصدار سابق (يبقى الأصل كما هو)
            old = os.path.join(base, legacy)
            if os.path.isdir(old):
                try:
                    shutil.copytree(old, d)
                except OSError:
                    pass
                break
    os.makedirs(d, exist_ok=True)
    return d


# ---------------------------------------------------------------- الإعدادات
def load_settings():
    try:
        with open(os.path.join(data_dir(), "settings.json"), encoding="utf-8") as fh:
            return json.load(fh)
    except (OSError, ValueError):
        return {}


def save_settings(settings):
    try:
        with open(os.path.join(data_dir(), "settings.json"), "w", encoding="utf-8") as fh:
            json.dump(settings, fh, ensure_ascii=False, indent=2)
    except OSError:
        pass


# ---------------------------------------------------------------- التحديثات
def parse_version(v):
    nums = [int(n) for n in re.findall(r"\d+", str(v))[:3]]
    return tuple(nums + [0] * (3 - len(nums)))


def update_url():
    return (load_settings().get("update_url") or UPDATE_URL).strip()


def check_for_update(url=None, timeout=6):
    """
    يرجع dict: newer (bool)، version، url (صفحة التنزيل https فقط)، notes.
    لا يرسل أي بيانات غير اسم البرنامج وإصداره في User-Agent.
    """
    url = (url or update_url()).strip()
    if not url:
        raise ValueError(tr("لم يتم ضبط رابط التحديثات بعد"))
    if not url.lower().startswith("https://"):
        raise ValueError(tr("رابط التحديثات يجب أن يبدأ بـ https://"))
    req = urllib.request.Request(url, headers={"User-Agent": f"{APP_NAME}/{APP_VERSION}",
                                               "Accept": "application/json"})
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        data = json.loads(resp.read(512 * 1024).decode("utf-8"))
    if "tag_name" in data:  # GitHub Releases
        version, page, notes = data["tag_name"], data.get("html_url", ""), data.get("body") or ""
    else:
        version, page, notes = data["version"], data.get("url", ""), data.get("notes") or ""
    if not str(page).lower().startswith("https://"):
        page = ""
    return {"newer": parse_version(version) > parse_version(APP_VERSION),
            "version": str(version).lstrip("vV"), "url": page, "notes": str(notes)[:2000]}


# ---------------------------------------------------------------- مكتبة التواقيع
LIBRARY_KINDS = {"signature": tr("توقيع"), "initials": tr("أحرف أولى"), "stamp": tr("ختم")}


def library_dir():
    d = os.path.join(data_dir(), "library")
    os.makedirs(d, exist_ok=True)
    return d


def library_list():
    """[(path, kind)] الأحدث أولًا."""
    items = []
    for name in os.listdir(library_dir()):
        if name.lower().endswith(".png") and "__" in name:
            kind = name.split("__", 1)[0]
            if kind in LIBRARY_KINDS:
                items.append((os.path.join(library_dir(), name), kind))
    items.sort(key=lambda t: os.path.getmtime(t[0]), reverse=True)
    return items


def library_save(img, kind):
    if kind not in LIBRARY_KINDS:
        raise ValueError(tr("نوع غير معروف"))
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    path = os.path.join(library_dir(), f"{kind}__{stamp}.png")
    img.save(path, format="PNG")
    return path


def library_delete(path):
    if os.path.dirname(os.path.abspath(path)) == os.path.abspath(library_dir()):
        os.remove(path)


# ---------------------------------------------------------------- صور التوقيع والنصوص
def remove_white_bg(img, threshold=235):
    """يحوّل الخلفية البيضاء/الفاتحة لصورة التوقيع إلى شفافة."""
    img = img.convert("RGBA")
    if img.width > 900:
        h = int(img.height * 900 / img.width)
        img = img.resize((900, h), Image.LANCZOS)
    px = [
        (r, g, b, 0) if (r > threshold and g > threshold and b > threshold) else (r, g, b, a)
        for (r, g, b, a) in img.getdata()
    ]
    img.putdata(px)
    bbox = img.getbbox()
    return img.crop(bbox) if bbox else img


def image_to_png_bytes(img):
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _fonts_archive():
    windir = os.environ.get("WINDIR", r"C:\Windows")
    fonts = os.path.join(windir, "Fonts")
    return fitz.Archive(fonts) if os.path.isdir(fonts) else None


# خطوط ويندوز التي تدعم العربية والإنجليزية: (الملف العادي، الملف العريض)
TEXT_FONTS = {
    "Tahoma": ("tahoma.ttf", "tahomabd.ttf"),
    "Arial": ("arial.ttf", "arialbd.ttf"),
    "Times New Roman": ("times.ttf", "timesbd.ttf"),
    "Courier New": ("cour.ttf", "courbd.ttf"),
}


def _text_css(px, color, bold, font="Tahoma", align="center", archive=None):
    css, family = "", "sans-serif"
    if archive is not None:
        regular, boldf = TEXT_FONTS.get(font, TEXT_FONTS["Tahoma"])
        family = "tf"
        css = (f"@font-face{{font-family:tf;src:url({regular});}}"
               f"@font-face{{font-family:tf;src:url({boldf});font-weight:bold;}}")
    r, g, b = color
    return css + (f"*{{font-family:{family};font-size:{px}px;color:rgb({r},{g},{b});"
                  f"text-align:{align};line-height:1.2;margin:0;{'font-weight:bold;' if bold else ''}}}")


def _text_html(text):
    return "<p dir='auto'>" + "<br>".join(html.escape(line) or "&nbsp;" for line in text.splitlines()) + "</p>"


def render_text_image(text, color=(10, 10, 10), bold=False, px=64, font="Tahoma", align="center"):
    """
    يرسم نصًا (عربي/إنجليزي، عدة أسطر) كصورة RGBA شفافة.
    يعتمد على محرك HTML في PyMuPDF لتشكيل الحروف العربية واتجاه RTL بشكل صحيح.
    """
    text = (text or "").strip()
    if not text:
        raise ValueError(tr("النص فارغ"))
    archive = _fonts_archive()
    css = _text_css(px, color, bold, font, align, archive)
    body = _text_html(text)
    lines = max(1, len(text.splitlines()))
    width = max(400, min(6000, int(px * 0.75 * max(len(t) for t in text.splitlines()) + px * 2)))
    height = int(px * 1.6 * lines + px)
    doc = fitz.open()
    page = doc.new_page(width=width, height=height)
    spare, _scale = page.insert_htmlbox(fitz.Rect(0, 0, width, height), body, css=css, archive=archive)
    if spare < 0:
        doc.close()
        raise ValueError(tr("النص طويل جدًا"))
    pix = page.get_pixmap(alpha=True)
    doc.close()
    img = Image.frombytes("RGBA", (pix.width, pix.height), pix.samples)
    bbox = img.getbbox()
    if not bbox:
        raise ValueError(tr("تعذر رسم النص"))
    pad = max(2, px // 10)
    x0, y0, x1, y1 = bbox
    return img.crop((max(0, x0 - pad), max(0, y0 - pad), min(img.width, x1 + pad), min(img.height, y1 + pad)))


def today_text():
    return datetime.date.today().strftime("%Y/%m/%d")


def with_opacity(img, opacity):
    if opacity >= 0.999:
        return img
    img = img.copy()
    a = img.getchannel("A").point(lambda v: int(v * opacity))
    img.putalpha(a)
    return img


# ---------------------------------------------------------------- التوقيع
SCOPES = ("page", "all", "first", "last")


def target_pages(scope, page_index, n):
    if scope == "all":
        return list(range(n))
    if scope == "first":
        return [0]
    if scope == "last":
        return [n - 1]
    return [page_index] if 0 <= page_index < n else []


def _insert_on_page(page, png, ratio, center, width_pts):
    """center بنقاط الصفحة كما تظهر على الشاشة (بعد التدوير)."""
    w, h = width_pts, width_pts * ratio
    pt = fitz.Point(*center)
    rot = page.rotation
    if rot:
        pt = pt * page.derotation_matrix
        if rot in (90, 270):
            w, h = h, w
    rect = fitz.Rect(pt.x - w / 2, pt.y - h / 2, pt.x + w / 2, pt.y + h / 2)
    page.insert_image(rect, stream=png, rotate=rot, overlay=True)


def stamp_items(pdf_path, out_path, items, progress=None):
    """
    يضع عدة عناصر (توقيع/أحرف أولى/ختم/نص/تاريخ) على PDF ويحفظه كملف جديد.
    كل عنصر dict فيه:
      img: صورة PIL RGBA
      rel: (x, y) مركز العنصر كنسبة من عرض/ارتفاع الصفحة الظاهرة (0..1)
      width: العرض بنقاط PDF
      scope: page | all | first | last
      page: رقم الصفحة (من صفر) عند scope == page
      opacity: (اختياري) 0..1
    """
    if not items:
        raise ValueError(tr("لا توجد عناصر للتوقيع"))
    prepared = []
    for it in items:
        img = with_opacity(it["img"], it.get("opacity", 1.0))
        prepared.append((it, image_to_png_bytes(img), img.height / img.width))
    doc = open_pdf(pdf_path)
    try:
        n = len(doc)
        jobs = []
        for it, png, ratio in prepared:
            for p in target_pages(it.get("scope", "page"), it.get("page", 0), n):
                jobs.append((p, it, png, ratio))
        if not jobs:
            raise ValueError(tr("لا توجد صفحات مطابقة في الملف"))
        for k, (p, it, png, ratio) in enumerate(jobs):
            page = doc[p]
            r = page.rect
            center = (it["rel"][0] * r.width, it["rel"][1] * r.height)
            _insert_on_page(page, png, ratio, center, it["width"])
            _tick(progress, k + 1, len(jobs))
        _save(doc, out_path)
    finally:
        doc.close()
    return out_path


def stamp_signature(pdf_path, out_path, sig_img, center_pts, width_pts, mode="current", current_page=0):
    """واجهة قديمة: توقيع واحد بإحداثيات بالنقاط."""
    with open_pdf(pdf_path) as d:
        r = d[current_page if mode == "current" else 0].rect
    scope = {"current": "page"}.get(mode, mode)
    item = dict(img=sig_img, rel=(center_pts[0] / r.width, center_pts[1] / r.height),
                width=width_pts, scope=scope, page=current_page)
    return stamp_items(pdf_path, out_path, [item])


def batch_stamp(paths, out_dir, items, progress=None):
    outs = []
    for i, p in enumerate(paths):
        name = os.path.splitext(os.path.basename(p))[0]
        od = out_dir or os.path.dirname(p)
        out = unique_path(os.path.join(od, tr("{0}_موقّع.pdf", name)))
        stamp_items(p, out, items)
        outs.append(out)
        _tick(progress, i + 1, len(paths))
    return outs


# ---------------------------------------------------------------- تعديل محتوى PDF
# كل الإحداثيات بنقاط الصفحة كما تظهر على الشاشة (بعد التدوير)، وتُحوَّل هنا لإحداثيات PDF.
EDIT_COLORS = {
    "أسود": (0, 0, 0), "أزرق": (20, 60, 180), "أحمر": (200, 30, 30), "أخضر": (20, 130, 60),
    "أصفر": (255, 214, 0), "برتقالي": (240, 130, 20), "رمادي": (110, 110, 110), "أبيض": (255, 255, 255),
}


def _to_pdf_rect(page, rect):
    r = fitz.Rect(rect)
    return (r * page.derotation_matrix).normalize() if page.rotation else r


def _to_pdf_point(page, pt):
    p = fitz.Point(pt)
    return p * page.derotation_matrix if page.rotation else p


def _to_view_rect(page, rect):
    r = fitz.Rect(rect)
    return (r * page.rotation_matrix).normalize() if page.rotation else r


def _rgb01(c):
    return tuple(v / 255 for v in c)


def text_line_at(page, point):
    """السطر النصي تحت النقطة (بإحداثيات العرض). يرجع dict أو None."""
    p = _to_pdf_point(page, point)
    best = None
    for block in page.get_text("dict", flags=fitz.TEXT_PRESERVE_WHITESPACE)["blocks"]:
        for line in block.get("lines", []):
            r = fitz.Rect(line["bbox"])
            if r.contains(p) or fitz.Rect(r.x0 - 2, r.y0 - 2, r.x1 + 2, r.y1 + 2).contains(p):
                spans = [s for s in line["spans"] if s["text"].strip()]
                if not spans:
                    continue
                best = (line, spans)
                break
        if best:
            break
    if not best:
        return None
    line, spans = best
    first = max(spans, key=lambda s: len(s["text"]))
    text = "".join(s["text"] for s in line["spans"]).replace("\xa0", " ").strip()
    c = first["color"]
    return {
        "rect": tuple(_to_view_rect(page, line["bbox"])),
        "text": _logical_arabic(text),
        "size": round(first["size"], 1),
        "color": ((c >> 16) & 255, (c >> 8) & 255, c & 255),
        "bold": bool(first["flags"] & 16) or "bold" in first["font"].lower(),
        "font": first["font"],
    }


def words_in_rect(page, rect):
    """مستطيلات الكلمات (بإحداثيات العرض) التي تتقاطع مع المستطيل المحدد — للتظليل والتسطير."""
    r = _to_pdf_rect(page, rect)
    out = []
    for w in page.get_text("words"):
        wr = fitz.Rect(w[:4])
        if wr.intersects(r):
            out.append(tuple(_to_view_rect(page, wr)))
    return out


def _insert_text(page, rect, text, size, color, bold, font, archive, align=None):
    if align is None:
        align = "right" if _has_arabic(text) else "left"
    css = _text_css(size, color, bold, font, align, archive)
    r = _to_pdf_rect(page, rect)
    page.insert_htmlbox(r, _text_html(text), css=css, archive=archive, scale_low=0, rotate=page.rotation)


def apply_edits(pdf_path, out_path, edits, progress=None):
    """
    يطبّق تعديلات المستخدم ويحفظ ملفًا جديدًا. كل تعديل dict فيه page و type وبيانات حسب النوع:
      text: rect, text, size, color, bold, font        replace: rect, text, size, color, bold, font, bg
      image: rect, png                                  highlight / underline / strike: rect, color
      ink: paths, color, width                          rect / ellipse: rect, color, width, fill
      line / arrow: p1, p2, color, width                whiteout: rect        redact: rect
      note: point, text, color
    """
    if not edits:
        raise ValueError(tr("لا توجد تعديلات للحفظ"))
    archive = _fonts_archive()
    doc = open_pdf(pdf_path)
    try:
        pages = sorted({e["page"] for e in edits})
        for k, pno in enumerate(pages):
            page = doc[pno]
            mine = [e for e in edits if e["page"] == pno]
            # 1) استبدال النص: إخفاء الأصل بخلفية ثم الكتابة فوقه (بدون المساس بالصور والخطوط)
            reps = [e for e in mine if e["type"] == "replace"]
            for e in reps:
                page.add_redact_annot(_to_pdf_rect(page, e["rect"]), fill=_rgb01(e.get("bg", (255, 255, 255))))
            if reps:
                page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE, graphics=fitz.PDF_REDACT_LINE_ART_NONE)
            # 2) الحجب النهائي: يحذف النص والصور تحت المنطقة فعليًا
            reds = [e for e in mine if e["type"] == "redact"]
            for e in reds:
                page.add_redact_annot(_to_pdf_rect(page, e["rect"]), fill=(0, 0, 0))
            if reds:
                page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_PIXELS)
            # 3) المحتوى الجديد والتعليقات بالترتيب الذي أضافه المستخدم
            for e in mine:
                t = e["type"]
                col = e.get("color", (0, 0, 0))
                if t in ("text", "replace"):
                    r = fitz.Rect(e.get("trect", e["rect"]))
                    # هامش حول النص حتى لا يُصغَّره محرك HTML ليتسع داخل المستطيل
                    pad = e["size"] * 0.6
                    r = fitz.Rect(r.x0 - pad, r.y0 - pad * 0.2, r.x1 + pad, r.y1 + e["size"])
                    if t == "replace":
                        r.y0 -= e["size"] * 0.15
                    _insert_text(page, r, e["text"], e["size"], col, e.get("bold", False),
                                 e.get("font", "Tahoma"), archive)
                elif t == "image":
                    page.insert_image(_to_pdf_rect(page, e["rect"]), stream=e["png"], rotate=page.rotation,
                                      keep_proportion=True, overlay=True)
                elif t == "whiteout":
                    page.draw_rect(_to_pdf_rect(page, e["rect"]), color=None, fill=(1, 1, 1), overlay=True)
                elif t in ("highlight", "underline", "strike"):
                    quads = [fitz.Rect(_to_pdf_rect(page, w)).quad for w in words_in_rect(page, e["rect"])]
                    target = quads or [_to_pdf_rect(page, e["rect"]).quad]
                    annot = {"highlight": page.add_highlight_annot, "underline": page.add_underline_annot,
                             "strike": page.add_strikeout_annot}[t](target)
                    annot.set_colors(stroke=_rgb01(col))
                    annot.update()
                elif t == "ink":
                    paths = [[tuple(_to_pdf_point(page, p)) for p in path] for path in e["paths"] if len(path) > 1]
                    if paths:
                        annot = page.add_ink_annot(paths)
                        annot.set_border(width=e.get("width", 2))
                        annot.set_colors(stroke=_rgb01(col))
                        annot.update()
                elif t in ("rect", "ellipse"):
                    r = _to_pdf_rect(page, e["rect"])
                    annot = page.add_rect_annot(r) if t == "rect" else page.add_circle_annot(r)
                    annot.set_border(width=e.get("width", 2))
                    fill = e.get("fill")
                    annot.set_colors(stroke=_rgb01(col), fill=_rgb01(fill) if fill else None)
                    annot.update()
                elif t in ("line", "arrow"):
                    annot = page.add_line_annot(_to_pdf_point(page, e["p1"]), _to_pdf_point(page, e["p2"]))
                    annot.set_border(width=e.get("width", 2))
                    annot.set_colors(stroke=_rgb01(col), fill=_rgb01(col))
                    if t == "arrow":
                        annot.set_line_ends(fitz.PDF_ANNOT_LE_NONE, fitz.PDF_ANNOT_LE_CLOSED_ARROW)
                    annot.update()
                elif t == "note":
                    annot = page.add_text_annot(_to_pdf_point(page, e["point"]), e["text"], icon="Comment")
                    annot.set_colors(stroke=_rgb01(col))
                    annot.update()
            _tick(progress, k + 1, len(pages))
        _save(doc, out_path)
    finally:
        doc.close()
    return out_path


# ---------------------------------------------------------------- أدوات PDF
def merge_pdfs(paths, out_path, progress=None):
    out = fitz.open()
    try:
        for i, p in enumerate(paths):
            with open_pdf(p) as d:
                out.insert_pdf(d)
            _tick(progress, i + 1, len(paths))
        _save(out, out_path)
    finally:
        out.close()
    return out_path


def split_pdf(path, out_dir, ranges_text="", progress=None):
    """بدون نطاق: كل صفحة في ملف. مع نطاق: ملف واحد يحتوي الصفحات المحددة."""
    name = os.path.splitext(os.path.basename(path))[0]
    outs = []
    with open_pdf(path) as src:
        if ranges_text.strip():
            pages = parse_ranges(ranges_text, len(src))
            d = fitz.open()
            for p in pages:
                d.insert_pdf(src, from_page=p, to_page=p)
            o = unique_path(os.path.join(out_dir, tr("{0}_مقتطف.pdf", name)))
            _save(d, o)
            d.close()
            outs.append(o)
            _tick(progress, 1, 1)
        else:
            for i in range(len(src)):
                d = fitz.open()
                d.insert_pdf(src, from_page=i, to_page=i)
                o = unique_path(os.path.join(out_dir, tr("{0}_صفحة_{1}.pdf", name, i + 1)))
                _save(d, o)
                d.close()
                outs.append(o)
                _tick(progress, i + 1, len(src))
    return outs


COMPRESS_LEVELS = {
    # (عتبة الدقة، الدقة المستهدفة، جودة JPEG) — None = بدون إعادة ضغط الصور
    "low": None,
    "medium": (200, 150, 75),
    "high": (120, 96, 55),
}


def compress_pdf(path, out_path, level="medium"):
    with open_pdf(path) as d:
        params = COMPRESS_LEVELS.get(level)
        if params:
            threshold, target, quality = params
            d.rewrite_images(dpi_threshold=threshold, dpi_target=target, quality=quality,
                             lossy=True, lossless=True, bitonal=True, color=True, gray=True)
        d.save(out_path, garbage=4, deflate=True, deflate_images=True, deflate_fonts=True,
               clean=True, use_objstms=True)
    return out_path


def rotate_pages(path, out_path, angle, ranges_text=""):
    with open_pdf(path) as d:
        pages = parse_ranges(ranges_text, len(d)) if ranges_text.strip() else range(len(d))
        for i in set(pages):
            d[i].set_rotation((d[i].rotation + angle) % 360)
        _save(d, out_path)
    return out_path


def delete_pages(path, out_path, ranges_text):
    with open_pdf(path) as d:
        drop = set(parse_ranges(ranges_text, len(d)))
        keep = [i for i in range(len(d)) if i not in drop]
        if not keep:
            raise ValueError(tr("لا يمكن حذف كل الصفحات"))
        d.select(keep)
        _save(d, out_path)
    return out_path


def reorder_pages(path, out_path, order_text):
    """order_text مثل '3,1,2' أو '5-1' (يمكن تكرار صفحة)."""
    with open_pdf(path) as d:
        d.select(parse_ranges(order_text, len(d)))
        _save(d, out_path)
    return out_path


def organize_pages(path, out_path, layout):
    """layout: [(رقم الصفحة الأصلي من صفر، تدوير إضافي بالدرجات)] بالترتيب الجديد."""
    if not layout:
        raise ValueError(tr("لا يمكن حفظ ملف بدون صفحات"))
    with open_pdf(path) as d:
        d.select([i for i, _ in layout])
        for page, (_, rot) in zip(d, layout):
            if rot % 360:
                page.set_rotation((page.rotation + rot) % 360)
        _save(d, out_path)
    return out_path


def page_count(path):
    with open_pdf(path) as d:
        return len(d)


def protect_pdf(path, out_path, password):
    if not password:
        raise ValueError(tr("أدخل كلمة المرور"))
    with open_pdf(path) as d:
        perm = (fitz.PDF_PERM_PRINT | fitz.PDF_PERM_COPY | fitz.PDF_PERM_ACCESSIBILITY
                | fitz.PDF_PERM_PRINT_HQ)
        d.save(out_path, garbage=3, deflate=True, encryption=fitz.PDF_ENCRYPT_AES_256,
               user_pw=password, owner_pw=password, permissions=perm)
    return out_path


def unlock_pdf(path, out_path, password):
    d = fitz.open(path)
    try:
        if d.needs_pass and not d.authenticate(password or ""):
            raise ValueError(tr("كلمة المرور غير صحيحة"))
        d.save(out_path, garbage=3, deflate=True, encryption=fitz.PDF_ENCRYPT_NONE)
    finally:
        d.close()
    return out_path


def watermark_image(text, color=(200, 30, 30), opacity=0.25, angle=45):
    img = render_text_image(text, color=color, bold=True, px=96)
    img = img.rotate(angle, expand=True, resample=Image.BICUBIC)
    return with_opacity(img, opacity)


def add_watermark(path, out_path, text, color=(200, 30, 30), opacity=0.25, angle=45, progress=None):
    img = watermark_image(text, color, opacity, angle)
    png = image_to_png_bytes(img)
    ratio = img.height / img.width
    with open_pdf(path) as d:
        for i, page in enumerate(d):
            r = page.rect
            width = min(r.width * 0.8, r.height * 0.8 / ratio)
            _insert_on_page(page, png, ratio, (r.width / 2, r.height / 2), width)
            _tick(progress, i + 1, len(d))
        _save(d, out_path)
    return out_path


# ---------------------------------------------------------------- معالجة متوازية (عمليات منفصلة)
# OCR وتحويل PDF إلى Word عمليات حسابية بطيئة لكل صفحة؛ نوزّع الصفحات على عدة عمليات (لا خيوط:
# PyMuPDF غير آمن مع الخيوط، والعمليات المنفصلة تستغل كل أنوية المعالج). كل عملية تفتح الملف بنفسها.
def worker_count(n_tasks, cap=6):
    cpu = os.cpu_count() or 2
    return max(1, min(n_tasks, cpu - 1, cap))


def _worker_init():
    import logging
    logging.disable(logging.INFO)
    try:  # أولوية أقل من الواجهة حتى يبقى البرنامج سلسًا أثناء المعالجة
        import win32api
        import win32process
        win32process.SetPriorityClass(win32api.GetCurrentProcess(), win32process.BELOW_NORMAL_PRIORITY_CLASS)
    except Exception:  # noqa: BLE001
        pass


def run_parallel(func, tasks, progress=None, weights=None, base=0.0, span=1.0):
    """
    يشغّل func(task) لكل مهمة في عمليات منفصلة ويرجع النتائج (بأي ترتيب).
    التقدم = base + span × (وزن المنجز / الوزن الكلي) من 1. الإلغاء يُنهي العمليات فورًا.
    """
    import multiprocessing as mp
    weights = weights or [1] * len(tasks)
    total_w = float(sum(weights)) or 1.0

    def tick(done_w):
        _tick(progress, base + span * done_w / total_w, 1.0)  # يرفع Cancelled عند الإلغاء

    n = worker_count(len(tasks))
    results, done_w = [], 0.0
    if n <= 1:
        for task, w in zip(tasks, weights):
            tick(done_w)
            results.append(func(task))
            done_w += w
            tick(done_w)
        return results
    pool = mp.get_context("spawn").Pool(n, initializer=_worker_init)
    try:
        it = pool.imap_unordered(_indexed_call, [(func, k, t) for k, t in enumerate(tasks)])
        while len(results) < len(tasks):
            try:
                k, res = it.next(timeout=0.4)
            except mp.TimeoutError:
                tick(done_w)  # فحص الإلغاء أثناء الانتظار
                continue
            results.append(res)
            done_w += weights[k]
            tick(done_w)
        pool.close()
        return results
    except BaseException:
        pool.terminate()
        raise
    finally:
        pool.join()


def _indexed_call(args):
    func, k, task = args
    return k, func(task)


# ---------------------------------------------------------------- التحويل
def images_to_pdf(paths, out_path, progress=None):
    doc = fitz.open()
    try:
        for i, p in enumerate(paths):
            with fitz.open(p) as img:
                pdf_bytes = img.convert_to_pdf()
            with fitz.open("pdf", pdf_bytes) as ipdf:
                doc.insert_pdf(ipdf)
            _tick(progress, i + 1, len(paths))
        _save(doc, out_path)
    finally:
        doc.close()
    return out_path


def pdf_to_images(path, out_dir, dpi=150, fmt="png", jpg_quality=85, progress=None):
    name = os.path.splitext(os.path.basename(path))[0]
    outs = []
    with open_pdf(path) as d:
        for i, page in enumerate(d):
            pix = page.get_pixmap(dpi=dpi, alpha=False)
            o = unique_path(os.path.join(out_dir, tr("{0}_صفحة_{1}.{2}", name, i + 1, fmt)))
            if fmt == "jpg":
                pix.save(o, jpg_quality=jpg_quality)
            else:
                pix.save(o)
            outs.append(o)
            _tick(progress, i + 1, len(d))
    return outs


def _w_pdf2docx_chunk(task):
    """(في عملية منفصلة) يحوّل مجموعة صفحات متتالية إلى ملف docx مستقل."""
    src, out, start, end = task
    from pdf2docx import Converter
    cv = Converter(src)
    try:
        cv.convert(out, start=start, end=end)
    finally:
        cv.close()
    return start, out


class _PdfDocxProgress:
    """يقرأ سجل pdf2docx «(i/n) Page» ليعطي تقدمًا لكل صفحة في التحويل غير المتوازي (ويُتيح الإلغاء)."""

    def __init__(self, progress):
        import logging
        import re
        self.progress, self.phase = progress, 0
        self.rx = re.compile(r"\((\d+)/(\d+)\) Page")
        outer = self

        class H(logging.Handler):
            def emit(self, record):
                msg = record.getMessage()
                if "[4/4]" in msg:
                    outer.phase = 1
                m = outer.rx.match(msg)
                if m:
                    i, n = int(m.group(1)), int(m.group(2))
                    # التحليل ~ النصف الأول والإنشاء ~ النصف الثاني (من القياس)
                    _tick(outer.progress, (outer.phase + (i - 1) / n) / 2, 1)  # يرفع Cancelled عند الإلغاء

        self.handler, self.logging = H(level=logging.INFO), logging

    def __enter__(self):
        root = self.logging.getLogger()
        self.old_level = root.level
        root.setLevel(self.logging.INFO)
        root.addHandler(self.handler)
        return self

    def __exit__(self, *exc):
        root = self.logging.getLogger()
        root.removeHandler(self.handler)
        root.setLevel(self.old_level)
        return False


def _merge_docx(parts, out):
    """يدمج ملفات docx بالترتيب مع فاصل قسم (صفحة جديدة) بين كل جزء والذي يليه."""
    import sys as _sys
    import types
    from copy import deepcopy
    try:
        import babel.dates  # noqa: F401
    except ImportError:  # babel مستبعدة من نسخة EXE (32 م.ب) — docxcompose تستخدمها فقط لتواريخ خصائص المستند
        stub = types.ModuleType("babel.dates")
        stub.format_datetime = lambda dt, *a, **k: dt.isoformat() if hasattr(dt, "isoformat") else str(dt)
        _sys.modules.setdefault("babel", types.ModuleType("babel"))
        _sys.modules["babel.dates"] = stub
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docxcompose.composer import Composer

    def end_with_section_break(doc):
        body = doc.element.body
        sect = body.find(qn("w:sectPr"))
        last = sect.getprevious()
        if last is not None and last.tag == qn("w:p") and last.find(qn("w:pPr") + "/" + qn("w:sectPr")) is None:
            ppr = last.find(qn("w:pPr"))
            if ppr is None:
                ppr = OxmlElement("w:pPr")
                last.insert(0, ppr)
            ppr.append(deepcopy(sect))
        else:
            para, ppr = OxmlElement("w:p"), OxmlElement("w:pPr")
            ppr.append(deepcopy(sect))
            para.append(ppr)
            sect.addprevious(para)

    composer = Composer(Document(parts[0]))
    for p in parts[1:]:
        end_with_section_break(composer.doc)
        composer.append(Document(p))
    composer.save(out)


def pdf_to_word(path, out_dir, progress=None):
    """
    PDF إلى Word. الملفات الصغيرة تُحوَّل مباشرة مع تقدم لكل صفحة، والكبيرة تُقسَّم إلى أجزاء
    تُحوَّل بالتوازي على عدة أنوية ثم تُدمج (أسرع ~4 مرات على جهاز بـ 8 أنوية).
    """
    import tempfile
    name = os.path.splitext(os.path.basename(path))[0]
    out = unique_path(os.path.join(out_dir, f"{name}.docx"))
    with open_pdf(path) as d:
        n = len(d)
    workers = worker_count(n)
    if n <= 1 or workers <= 1:
        from pdf2docx import Converter
        cv = Converter(path)
        try:
            with _PdfDocxProgress(progress):
                cv.convert(out)
        finally:
            cv.close()
        _tick(progress, 1, 1)
        return out
    size = max(1, min(10, -(-n // (workers * 2))))  # أجزاء صغيرة = تقدم أدق وتوزيع أعدل
    tmp = tempfile.mkdtemp(prefix="pdfya_docx_")
    try:
        chunks = [(path, os.path.join(tmp, f"part_{s:05d}.docx"), s, min(s + size, n)) for s in range(0, n, size)]
        done = dict(run_parallel(_w_pdf2docx_chunk, chunks, progress, weights=[c[3] - c[2] for c in chunks],
                                 span=0.9))
        _tick(progress, 0.92, 1)
        _merge_docx([done[k] for k in sorted(done)], out)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    _tick(progress, 1, 1)
    return out


def _has_arabic(text):
    return any("؀" <= ch <= "ۿ" for ch in text)


def _logical_arabic(text):
    """
    بعض ملفات PDF تخزّن العربي كحروف مُشكّلة (Presentation Forms) بترتيب بصري معكوس.
    نعيدها لحروف عادية بترتيب منطقي، مع الحفاظ على اتجاه الأرقام والكلمات اللاتينية.
    """
    if not any("ﭐ" <= ch <= "﷿" or "ﹰ" <= ch <= "﻿" for ch in text):
        return text
    text = unicodedata.normalize("NFKC", text)[::-1]
    out, run = [], []
    for ch in text:
        if _has_arabic(ch) or ch.isspace() or ch in "()[]{}<>":
            if run:
                out.append("".join(run)[::-1])
                run = []
            out.append({"(": ")", ")": "(", "[": "]", "]": "[", "{": "}", "}": "{", "<": ">", ">": "<"}.get(ch, ch))
        else:
            run.append(ch)
    if run:
        out.append("".join(run)[::-1])
    return "".join(out)


def _cell_value(val):
    """يحوّل الأرقام النصية إلى أرقام حتى تعمل معادلات Excel عليها."""
    txt = val.replace(",", "").replace("٬", "")
    if txt and txt.lstrip("-").replace(".", "", 1).isdigit():
        return float(txt) if "." in txt else int(txt)
    return val


def pdf_to_excel(path, out_dir, progress=None):
    """يستخرج الجداول من كل صفحة، كل جدول في ورقة Excel مستقلة."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    name = os.path.splitext(os.path.basename(path))[0]
    wb = Workbook()
    wb.remove(wb.active)
    count = 0
    with open_pdf(path) as d:
        for pno, page in enumerate(d):
            for tno, table in enumerate(page.find_tables().tables):
                rows = table.extract()
                if not rows:
                    continue
                count += 1
                ws = wb.create_sheet(tr("ص{0}-جدول{1}", pno + 1, tno + 1)[:31])
                arabic = False
                for r, row in enumerate(rows, 1):
                    for c, val in enumerate(row, 1):
                        val = _logical_arabic((val or "").strip())
                        arabic = arabic or _has_arabic(val)
                        cell = ws.cell(row=r, column=c, value=_cell_value(val))
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
                        if r == 1:
                            cell.font = Font(bold=True)
                for col in ws.columns:
                    width = max(len(str(c.value or "")) for c in col)
                    ws.column_dimensions[col[0].column_letter].width = min(60, max(8, width + 2))
                ws.sheet_view.rightToLeft = arabic
            _tick(progress, pno + 1, len(d))
    if not count:
        raise ValueError(tr("لم يتم العثور على جداول في: {0}", os.path.basename(path)))
    out = unique_path(os.path.join(out_dir, f"{name}.xlsx"))
    wb.save(out)
    return out


def pdf_to_pptx(path, out_dir, dpi=150, progress=None):
    """كل صفحة تصبح شريحة (صورة الصفحة بكامل الشريحة)."""
    from pptx import Presentation
    from pptx.util import Emu

    emu = 12700  # نقطة PDF = 12700 EMU
    name = os.path.splitext(os.path.basename(path))[0]
    prs = Presentation()
    blank = prs.slide_layouts[6]
    with open_pdf(path) as d:
        r = d[0].rect
        prs.slide_width, prs.slide_height = Emu(int(r.width * emu)), Emu(int(r.height * emu))
        sw, sh = prs.slide_width, prs.slide_height
        for i, page in enumerate(d):
            pix = page.get_pixmap(dpi=dpi, alpha=False)
            pw, ph = page.rect.width * emu, page.rect.height * emu
            scale = min(sw / pw, sh / ph)
            w, h = int(pw * scale), int(ph * scale)
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(io.BytesIO(pix.tobytes("png")), (sw - w) // 2, (sh - h) // 2, w, h)
            _tick(progress, i + 1, len(d))
    out = unique_path(os.path.join(out_dir, f"{name}.pptx"))
    prs.save(out)
    return out


# ---------------------------------------------------------------- OCR (Tesseract)
OCR_LANGS = {tr("عربي + إنجليزي"): "ara+eng", tr("عربي"): "ara", tr("إنجليزي"): "eng"}
TESSERACT_URL = "https://github.com/UB-Mannheim/tesseract/wiki"


def find_tessdata():
    """يبحث عن مجلد tessdata الخاص بـ Tesseract. يرجع المسار أو None."""
    cands = []
    if os.environ.get("TESSDATA_PREFIX"):
        cands.append(os.environ["TESSDATA_PREFIX"])
    base = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    cands += [
        os.path.join(base, "tesseract", "tessdata"),
        os.path.join(os.path.dirname(sys.executable), "tesseract", "tessdata"),
        r"C:\Program Files\Tesseract-OCR\tessdata",
        r"C:\Program Files (x86)\Tesseract-OCR\tessdata",
        os.path.join(os.environ.get("LOCALAPPDATA", ""), "Programs", "Tesseract-OCR", "tessdata"),
    ]
    exe = shutil.which("tesseract")
    if exe:
        cands.append(os.path.join(os.path.dirname(exe), "tessdata"))
    for c in cands:
        try:
            if c and os.path.isdir(c) and any(n.endswith(".traineddata") for n in os.listdir(c)):
                return c
        except OSError:
            pass
    return None


def _ocr_ready(lang):
    tessdata = find_tessdata()
    if not tessdata:
        raise RuntimeError(
            tr("التعرف على النص (OCR) يحتاج برنامج Tesseract المجاني. ثبّته من:\n{0}\nواختر اللغة العربية (Arabic) أثناء التثبيت، ثم أعد تشغيل البرنامج.", TESSERACT_URL)
        )
    missing = [x for x in lang.split("+") if not os.path.exists(os.path.join(tessdata, f"{x}.traineddata"))]
    if missing:
        raise RuntimeError(tr("ملفات اللغة غير مثبتة في Tesseract: {0}", ', '.join(missing)))
    return tessdata


def _open_any(path):
    """يفتح PDF أو صورة كمستند PyMuPDF."""
    if os.path.splitext(path)[1].lower() == ".pdf":
        return open_pdf(path)
    return fitz.open(path)


def _w_ocr_page(task):
    """(في عملية منفصلة) OCR لصفحة واحدة. mode=pdf يرجع PDF قابلًا للبحث، mode=text يرجع النص."""
    path, i, lang, dpi, tessdata, mode = task
    with _open_any(path) as src:
        page = src[i]
        if mode == "pdf":
            pix = page.get_pixmap(dpi=dpi, alpha=False)
            return i, pix.pdfocr_tobytes(language=lang, tessdata=tessdata)
        tp = page.get_textpage_ocr(language=lang, dpi=dpi, full=True, tessdata=tessdata)
        return i, page.get_text("text", textpage=tp).strip()  # بدون sort: الترتيب يعكس الكلمات العربية


def ocr_pdf(path, out_dir, lang="ara+eng", dpi=300, progress=None):
    """
    ينتج PDF قابلًا للبحث والنسخ من ملف ممسوح (PDF أو صورة).
    الصفحات التي تحتوي نصًا بالفعل تُنسخ كما هي، والباقي يُعالج بالتوازي على عدة أنوية.
    """
    tessdata = _ocr_ready(lang)
    name = os.path.splitext(os.path.basename(path))[0]
    out = unique_path(os.path.join(out_dir, tr("{0}_قابل_للبحث.pdf", name)))
    with _open_any(path) as src:
        n = len(src)
        need = [i for i in range(n) if not (src.is_pdf and src[i].get_text().strip())]
    ocr = dict(run_parallel(_w_ocr_page, [(path, i, lang, dpi, tessdata, "pdf") for i in need], progress, span=0.97))
    result = fitz.open()
    try:
        with _open_any(path) as src:
            for i in range(n):
                if i in ocr:
                    with fitz.open("pdf", ocr[i]) as op:
                        result.insert_pdf(op)
                else:
                    result.insert_pdf(src, from_page=i, to_page=i)
        _save(result, out)
    finally:
        result.close()
    _tick(progress, 1, 1)
    return out


def ocr_text(path, lang="ara+eng", dpi=300, progress=None):
    """يرجع قائمة بنص كل صفحة (الصفحات تُعالج بالتوازي)."""
    tessdata = _ocr_ready(lang)
    with _open_any(path) as src:
        n = len(src)
    texts = dict(run_parallel(_w_ocr_page, [(path, i, lang, dpi, tessdata, "text") for i in range(n)], progress))
    return [texts[i] for i in range(n)]


def ocr_to_txt(path, out_dir, lang="ara+eng", dpi=300, progress=None):
    name = os.path.splitext(os.path.basename(path))[0]
    pages = ocr_text(path, lang, dpi, progress)
    out = unique_path(os.path.join(out_dir, tr("{0}_نص.txt", name)))
    with open(out, "w", encoding="utf-8-sig") as fh:
        fh.write("\n\n".join(tr("===== صفحة {0} =====\n{1}", i + 1, t) for i, t in enumerate(pages)))
    return out


def ocr_to_word(path, out_dir, lang="ara+eng", dpi=300, progress=None):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    name = os.path.splitext(os.path.basename(path))[0]
    pages = ocr_text(path, lang, dpi, progress)
    doc = Document()
    for i, text in enumerate(pages):
        if i:
            doc.add_page_break()
        for line in text.splitlines() or [""]:
            p = doc.add_paragraph(line)
            if _has_arabic(line):
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                p._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
                for run in p.runs:
                    rpr = run._r.get_or_add_rPr()
                    rpr.append(OxmlElement("w:rtl"))
                    run.font.name = "Arial"
                    rpr.get_or_add_rFonts().set(qn("w:cs"), "Arial")
    out = unique_path(os.path.join(out_dir, f"{name}_OCR.docx"))
    doc.save(out)
    return out


def _find_soffice():
    for cand in ("soffice", "libreoffice"):
        p = shutil.which(cand)
        if p:
            return p
    for p in (
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ):
        if os.path.exists(p):
            return p
    return None


DUMMY_PASSWORD = "pdfya-no-password-9"  # الملفات غير المحمية تتجاهلها، والمحمية تفشل فورًا


class OfficeTimeout(RuntimeError):
    """توقف Office/LibreOffice عن الاستجابة وأوقفناه."""


class OfficeSession:
    """
    تحويل Word / Excel / PowerPoint إلى PDF عبر Microsoft Office، مع:
      - نسخة واحدة من كل برنامج لكل الملفات (يوفر 3-6 ثوانٍ تشغيل لكل ملف إضافي).
      - فتح صامت: بدون رسائل أو ماكرو أو تحديث روابط، والملفات المحمية بكلمة مرور تفشل فورًا بدل انتظار نافذة مخفية.
      - مراقِب (watchdog) في خيط منفصل: يُنهي عملية Office التي شغّلناها نحن فقط عند الإلغاء أو تجاوز المهلة.
      - تقدير زمن حسب عدد الصفحات واللغة، وتقدّم يتحرك أثناء التحويل.
    progress: دالة RunBar (done, total) ولها اختياريًا .cancelled() و .status(text).
    """
    EXES = {"word": "WINWORD.EXE", "excel": "EXCEL.EXE", "powerpoint": "POWERPNT.EXE"}
    OPEN_TIMEOUT = 120  # ثانية: فتح الملف (نافذة مخفية عالقة = توقف هنا)

    def __init__(self, progress=None):
        self.progress = progress
        self.apps = {}
        self.pids = {}
        self._killed = None  # None أو "cancel" أو "timeout"
        self._deadline = None
        self._estimate = None
        self._task = None  # (i, n, label) للعرض
        self._stop = threading.Event()
        self._lock = threading.Lock()

    # ---- دورة الحياة
    def __enter__(self):
        import pythoncom
        pythoncom.CoInitialize()
        self._watch = threading.Thread(target=self._watchdog, daemon=True)
        self._watch.start()
        return self

    def __exit__(self, *exc):
        self._stop.set()
        for kind, app in list(self.apps.items()):
            try:
                app.Quit()
            except Exception:  # noqa: BLE001
                pass
        self.apps.clear()
        time.sleep(0.3)
        self._kill_all()  # أي نسخة لم تُغلق بسلام
        import pythoncom
        pythoncom.CoUninitialize()
        return False

    # ---- العمليات
    @staticmethod
    def _running(exe):
        try:
            out = subprocess.run(["tasklist", "/FI", f"IMAGENAME eq {exe}", "/FO", "CSV", "/NH"],
                                 capture_output=True, text=True, timeout=15,
                                 creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0)).stdout
        except Exception:  # noqa: BLE001
            return set()
        pids = set()
        for line in out.splitlines():
            parts = [p.strip('"') for p in line.split('","')]
            if len(parts) > 1 and parts[1].isdigit():
                pids.add(int(parts[1]))
        return pids

    def _kill_all(self):
        import signal
        for pid in list(self.pids.values()):
            if pid:
                try:
                    os.kill(pid, signal.SIGTERM)  # على ويندوز = TerminateProcess
                except OSError:
                    pass
        self.pids.clear()

    def _app(self, kind):
        if kind in self.apps:
            return self.apps[kind]
        import win32com.client as wc
        exe = self.EXES[kind]
        before = self._running(exe)
        app = wc.DispatchEx({"word": "Word.Application", "excel": "Excel.Application",
                             "powerpoint": "PowerPoint.Application"}[kind])
        new = self._running(exe) - before
        self.pids[kind] = max(new) if new else None  # النسخة التي شغّلناها نحن فقط
        for attr, value in {
            "word": (("Visible", False), ("DisplayAlerts", 0), ("ScreenUpdating", False), ("AutomationSecurity", 3)),
            "excel": (("Visible", False), ("DisplayAlerts", False), ("AskToUpdateLinks", False),
                      ("ScreenUpdating", False), ("AutomationSecurity", 3)),
            "powerpoint": (("DisplayAlerts", 1), ("AutomationSecurity", 3)),
        }[kind]:
            try:
                setattr(app, attr, value)
            except Exception:  # noqa: BLE001
                pass
        self.apps[kind] = app
        return app

    # ---- المراقِب: إلغاء + مهلة + تقدّم متحرك
    def _watchdog(self):
        last_status = 0
        while not self._stop.wait(0.4):
            with self._lock:
                deadline, est, task, started = self._deadline, self._estimate, self._task, getattr(self, "_t0", None)
            if task is None:
                continue
            p = self.progress
            cancelled = bool(p and getattr(p, "cancelled", None) and p.cancelled())
            timed_out = deadline is not None and time.time() > deadline
            if cancelled or timed_out:
                self._killed = "cancel" if cancelled else "timeout"
                self.apps.clear()
                self._kill_all()
                continue
            if p and started:
                i, n, label = task
                elapsed = time.time() - started
                if est:
                    try:
                        p(i + min(0.95, elapsed / est), n)
                    except Exception:  # noqa: BLE001
                        pass
                if hasattr(p, "status") and time.time() - last_status >= 1:
                    last_status = time.time()
                    if est and elapsed <= est:
                        extra = tr("— مضى {0} ث من حوالي {1} ث", int(elapsed), int(est))
                    elif est:
                        extra = tr("— مضى {0} ث (يأخذ وقتًا أطول قليلًا من المتوقع)", int(elapsed))
                    else:
                        extra = tr("— مضى {0} ث", int(elapsed))
                    p.status(label + "  " + extra)

    def _begin(self, i, n, label, timeout):
        with self._lock:
            self._task = (i, n, label)
            self._t0 = time.time()
            self._estimate = None
            self._deadline = time.time() + timeout

    def _set_estimate(self, label, est):
        with self._lock:
            spent = time.time() - (getattr(self, "_t0", None) or time.time())
            self._estimate = spent + max(2.0, est)
            self._deadline = time.time() + 90 + est * 3
            self._task = (self._task[0], self._task[1], label)

    def _end(self):
        with self._lock:
            self._task = self._deadline = self._estimate = None

    def _raise_if_killed(self, err):
        if self._killed == "cancel":
            raise Cancelled()
        if self._killed == "timeout":
            raise OfficeTimeout(tr("توقف Microsoft Office عن الاستجابة أثناء تحويل الملف، فأوقفناه. "
                                  "جرّب فتح الملف في Office للتأكد أنه سليم وغير محمي، ثم أعد المحاولة."))
        msg = str(err)
        if "password" in msg.lower() or "كلمة" in msg:
            raise ValueError(tr("الملف محمي بكلمة مرور. افتحه في Office واحفظه بدون كلمة مرور، ثم حوّله."))
        raise err

    # ---- التحويل
    def to_pdf(self, src, out, i=0, n=1):
        ext = os.path.splitext(src)[1].lower()
        src, out = os.path.abspath(src), os.path.abspath(out)
        name = os.path.basename(src)
        kind = ("word" if ext in (".doc", ".docx", ".rtf", ".odt") else
                "excel" if ext in (".xls", ".xlsx", ".xlsm", ".csv") else
                "powerpoint" if ext in (".ppt", ".pptx") else None)
        if kind is None:
            raise ValueError(tr("صيغة غير مدعومة"))
        self._killed = None
        self._begin(i, n, tr("يفتح «{0}»...", name), self.OPEN_TIMEOUT)
        try:
            app = self._app(kind)
            if kind == "word":
                # كلمة مرور وهمية: الملف المحمي يفشل فورًا بدل نافذة مخفية تنتظر للأبد
                doc = app.Documents.Open(src, ConfirmConversions=False, ReadOnly=True, AddToRecentFiles=False,
                                         PasswordDocument=DUMMY_PASSWORD, Visible=False, NoEncodingDialog=True)
                try:
                    pages = doc.ComputeStatistics(2)
                    sample = doc.Range(0, min(doc.Content.End, 20000)).Text
                    arabic = _has_arabic(sample)
                    # من القياس على هذا الجهاز: العربي ~9 ث ثابتة + 1.8 ث/صفحة، والإنجليزي ~2 ث + 0.4 ث/صفحة
                    est = 9 + pages * 1.8 if arabic else 2 + pages * 0.4
                    self._set_estimate(tr("يحوّل «{0}» ({1} صفحة)", name, pages), est)
                    doc.ExportAsFixedFormat(out, 17, OpenAfterExport=False, OptimizeFor=0, IncludeDocProps=True,
                                            CreateBookmarks=1, DocStructureTags=True)
                finally:
                    try:
                        doc.Close(0)
                    except Exception:  # noqa: BLE001
                        pass
            elif kind == "excel":
                wb = app.Workbooks.Open(src, UpdateLinks=0, ReadOnly=True, Password=DUMMY_PASSWORD,
                                        IgnoreReadOnlyRecommended=True, AddToMru=False)
                try:
                    sheets = wb.Worksheets.Count
                    self._set_estimate(tr("يحوّل «{0}» ({1} ورقة)", name, sheets), 2 + sheets * 0.6)
                    wb.ExportAsFixedFormat(0, out)
                finally:
                    try:
                        wb.Close(False)
                    except Exception:  # noqa: BLE001
                        pass
            else:
                pres = app.Presentations.Open(src, ReadOnly=True, Untitled=False, WithWindow=False)
                try:
                    slides = pres.Slides.Count
                    self._set_estimate(tr("يحوّل «{0}» ({1} شريحة)", name, slides), 2 + slides * 0.3)
                    pres.SaveAs(out, 32)
                finally:
                    try:
                        pres.Close()
                    except Exception:  # noqa: BLE001
                        pass
        except Exception as e:  # noqa: BLE001
            self._raise_if_killed(e)
        finally:
            self._end()
        if not os.path.exists(out):
            raise RuntimeError(tr("لم يُنشئ Office ملف PDF لهذا الملف."))
        return out


def _libreoffice_to_pdf(soffice, src, out_dir, progress=None, timeout=600):
    """LibreOffice بملف إعدادات مستقل (حتى لا يتعارض مع نسخة مفتوحة عند المستخدم) مع إلغاء ومهلة."""
    profile = os.path.join(data_dir(), "lo_profile").replace("\\", "/")
    proc = subprocess.Popen(
        [soffice, f"-env:UserInstallation=file:///{profile}", "--headless", "--norestore",
         "--convert-to", "pdf", "--outdir", out_dir, src],
        stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    t0 = time.time()
    while proc.poll() is None:
        time.sleep(0.4)
        if progress and getattr(progress, "cancelled", None) and progress.cancelled():
            proc.kill()
            raise Cancelled()
        if time.time() - t0 > timeout:
            proc.kill()
            raise OfficeTimeout(tr("توقف LibreOffice عن الاستجابة أثناء التحويل، فأوقفناه."))
        if progress and hasattr(progress, "status"):
            progress.status(tr("يحوّل «{0}» عبر LibreOffice — مضى {1} ث", os.path.basename(src), int(time.time() - t0)))


def office_to_pdf(src, out_dir, session=None, progress=None, i=0, n=1):
    """Word/Excel/PowerPoint -> PDF. يستخدم Microsoft Office إن وُجد، وإلا LibreOffice."""
    name = os.path.splitext(os.path.basename(src))[0]
    out = unique_path(os.path.join(out_dir, f"{name}.pdf"))
    errors = []
    if sys.platform.startswith("win"):
        try:
            if session is not None:
                return session.to_pdf(src, out, i, n)
            with OfficeSession(progress) as s:
                return s.to_pdf(src, out, i, n)
        except (Cancelled, ValueError, OfficeTimeout):
            raise
        except Exception as e:  # noqa: BLE001
            errors.append(f"Office: {e}")
    soffice = _find_soffice()
    if soffice:
        try:
            _libreoffice_to_pdf(soffice, src, out_dir, progress)
            produced = os.path.join(out_dir, name + ".pdf")
            if os.path.exists(produced):
                if produced != out:
                    shutil.move(produced, out)
                return out
        except (Cancelled, OfficeTimeout):
            raise
        except Exception as e:  # noqa: BLE001
            errors.append(f"LibreOffice: {e}")
    raise RuntimeError(
        tr("لا يوجد Microsoft Office ولا LibreOffice على الجهاز لتنفيذ هذا التحويل.")
        + (("\n" + " | ".join(errors)) if errors else "")
    )
